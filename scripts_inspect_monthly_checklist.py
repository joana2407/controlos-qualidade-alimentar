from pathlib import Path
import json
from openpyxl import load_workbook

source = Path('/home/ubuntu/upload/pasted_file_ikNPfg_checklistmensal.xlsx')
workbook = load_workbook(source, data_only=True)
sheet = workbook.worksheets[0]
rows = [[str(value).strip() if value is not None else '' for value in row] for row in sheet.iter_rows(values_only=True)]
non_empty_rows = [index for index, row in enumerate(rows) if any(row)]
header_row_index = next(index for index in non_empty_rows if 'Armazém de Matéria-Prima' in ' '.join(rows[index]))
area_headers = {index: value for index, value in enumerate(rows[header_row_index]) if value and index >= 5}
items = []
for row in rows[header_row_index + 4:]:
    item = row[0].strip() if row else ''
    if not item:
        continue
    area_columns = [index for index, value in enumerate(row) if value.lower() == 'x' and index in area_headers]
    if area_columns:
        items.append({'point': item, 'areas': [area_headers[index] for index in area_columns]})
print(json.dumps({'sheet': sheet.title, 'header_row_index': header_row_index, 'areas': area_headers, 'items': items}, ensure_ascii=False, indent=2))
